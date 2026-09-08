from __future__ import annotations

from collections.abc import Iterator
from pathlib import Path
from typing import ClassVar

from pptx import Presentation

from universal_parser.core.router import register
from universal_parser.core.schema import Element, TableData
from universal_parser.core.sniffer import FileType
from universal_parser.extractors.base import BaseExtractor


@register
class PPTXExtractor(BaseExtractor):
    """
    Extractor for PowerPoint presentations (.pptx).
    Handles:
        - Slide titles -> Level 1 Headings with slide numbers
        - Text boxes and bullet points -> Paragraphs and List Items
        - Presentation tables -> Structured TableData & Markdown tables
    """

    supported_types: ClassVar[list[FileType]] = [FileType.PPTX]

    def stream(self, path: str | Path) -> Iterator[Element]:
        path_str = str(path)

        try:
            prs = Presentation(path_str)

        except Exception:  # noqa: BLE001
            return

        for slide_num, slide in enumerate(prs.slides, start=1):
            # 1. Slide Title (Heading 1)
            if slide.shapes.title and slide.shapes.title.text.strip():
                title = slide.shapes.title.text.strip()
                yield Element(
                    type="heading",
                    level=1,
                    text=f"Slide {slide_num}: {title}",
                    page=slide_num,
                    markdown_repr=f"# Slide {slide_num}: {title}",
                    confidence=1.0,
                )
            else:
                yield Element(
                    type="heading",
                    level=1,
                    text=f"Slide {slide_num}",
                    page=slide_num,
                    markdown_repr=f"# Slide {slide_num}",
                    confidence=1.0,
                )

            # 2. Iterate through shapes on the slide
            for shape in slide.shapes:
                # Skip title shape since already extracted
                if shape == slide.shapes.title:
                    continue
                # A. Handle Tables in slides
                if shape.has_table:
                    table = shape.table
                    rows = []
                    for row in table.rows:
                        row_cells = [cell.text.strip() for cell in row.cells]
                        if any(row_cells):
                            rows.append(row_cells)

                    if not rows:
                        continue

                    headers = rows[0]
                    data_rows = rows[1:]

                    md_header = "| " + " | ".join(headers) + " |"
                    md_separator = "| " + " | ".join(["---"] * len(headers)) + " |"
                    md_rows = ["| " + " | ".join(r) + " |" for r in data_rows]
                    markdown_repr = "\n".join([md_header, md_separator] + md_rows)

                    yield Element(
                        type="table",
                        page=slide_num,
                        text=f"Slide {slide_num} Table",
                        data=TableData(headers=headers, rows=data_rows),
                        markdown_repr=markdown_repr,
                        confidence=1.0,
                    )

                # B. Handle Text Frames (paragraphs and bullet points)
                elif shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if not text:
                            continue

                        is_bullet = paragraph.level > 0
                        elem_type = "list_item" if is_bullet else "paragraph"
                        prefix = "  " * paragraph.level + "- " if is_bullet else ""

                        yield Element(
                            type=elem_type,
                            page=slide_num,
                            text=text,
                            markdown_repr=f"{prefix} {text}",
                            confidence=1.0,
                        )
