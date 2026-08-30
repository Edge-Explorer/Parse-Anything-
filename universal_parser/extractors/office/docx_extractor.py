from __future__ import annotations

from collections.abc import Iterator
from pathlib import Path
from typing import ClassVar

from docx import Document as DocxDocument
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from universal_parser.core.router import register
from universal_parser.core.schema import Element, TableData
from universal_parser.core.sniffer import FileType
from universal_parser.extractors.base import BaseExtractor

# Map Word built-in heading style names to heading levels
_HEADING_STYLE_MAP: dict[str, int] = {
    "heading 1": 1,
    "heading 2": 2,
    "heading 3": 3,
    "heading 4": 4,
    "heading 5": 5,
    "heading 6": 6,
    "title": 1,  # Word's "Title" style treated as H1
    "subtitle": 2,  # Word's "Subtitle" style treated as H2
}


@register
class DocxExtractor(BaseExtractor):
    """
    Extractor for .docx files using python-docx.
    Handles:
        - Paragraphs and headings in correct document order
        - Tables (headers inferred from first row)
        - Preserves reading order by iterating document body directly

    Does NOT handle:
        - Legacy .doc binary format (needs olefile — Phase 6)
        - Embedded images/charts (Phase 6)
    """

    supported_types: ClassVar[list[FileType]] = [FileType.DOCX]

    def stream(self, path: str | Path) -> Iterator[Element]:
        """Stream elements from a DOCX file in document reading order."""
        try:
            doc = DocxDocument(str(path))
        except Exception:  # noqa: BLE001
            # Corrupted or password-protected DOCX — exit cleanly
            return

        try:
            # Iterate body children directly to preserve paragraph + table order.
            # doc.paragraphs and doc.tables are separate lists — using them
            # would give all paragraphs first then all tables, losing real order.
            for child in doc.element.body:
                tag = child.tag

                # Paragraph element
                if tag == qn("w:p"):
                    element = self._process_paragraph(Paragraph(child, doc))
                    if element is not None:
                        yield element

                # Table element
                elif tag == qn("w:tbl"):
                    element = self._process_table(Table(child, doc))
                    if element is not None:
                        yield element

        except Exception:  # noqa: BLE001
            # If anything fails mid-document, we stop cleanly.
            # We do not crash — whatever was yielded before the error is valid.
            return

    def _process_paragraph(self, para: Paragraph) -> Element | None:
        """Convert a python-docx Paragraph into an Element."""
        text = para.text.strip()

        # Skip empty paragraphs — they carry no information for RAG
        if not text:
            return None

        style_name = para.style.name.lower() if para.style else ""
        heading_level = _HEADING_STYLE_MAP.get(style_name)

        if heading_level is not None:
            return Element(
                type="heading",
                level=heading_level,
                text=text,
            )

        # Check for list items (Word list styles contain "list" in the name)
        if "list" in style_name:
            return Element(
                type="list_item",
                text=text,
            )
        return Element(
            type="paragraph",
            text=text,
        )

    def _process_table(self, table: Table) -> Element | None:
        """Convert a python-docx Table into an Element with TableData."""
        rows = []
        for row in table.rows:
            row_data = [cell.text.strip() for cell in row.cells]
            rows.append(row_data)

        if not rows:
            return None

        # Treat the first row as headers
        headers = rows[0]
        data_rows = rows[1:]

        # Build a markdown representation for quick use in RAG prompts
        md_header = "| " + " | ".join(headers) + " |"
        md_separator = "| " + " | ".join(["---"] * len(headers)) + " |"
        md_rows = ["| " + " | ".join(row) + " |" for row in data_rows]
        markdown_repr = "\n".join([md_header, md_separator] + md_rows)

        return Element(
            type="table",
            data=TableData(headers=headers, rows=data_rows),
            markdown_repr=markdown_repr,
        )
