from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from universal_parser.core.engine import parse
from universal_parser.core.schema import Document


@pytest.fixture
def temp_pptx_file(tmp_path: Path) -> Path:
    """Generate a sample .pptx file with text and tables."""
    file_path = tmp_path / "sample.pptx"

    prs = Presentation()

    # Slide 1: Title & Bullet points
    slide_layout = prs.slide_layouts[1]  # Title and Content layout
    slide1 = prs.slides.add_slide(slide_layout)
    slide1.shapes.title.text = "Q3 Architecture Review"

    tf = slide1.placeholders[1].text_frame
    tf.text = "Key System Achievements"
    p2 = tf.add_paragraph()
    p2.text = "Zero GPU footprint achieved"
    p2.level = 1

    # Slide 2: Table Slide
    slide2 = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
    table_shape = slide2.shapes.add_table(3, 2, Inches(1), Inches(1), Inches(6), Inches(2))
    table = table_shape.table
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Target"
    table.cell(1, 0).text = "Latency"
    table.cell(1, 1).text = "<100ms"
    table.cell(2, 0).text = "Memory"
    table.cell(2, 1).text = "<250MB"

    prs.save(str(file_path))
    return file_path


def test_pptx_returns_document(temp_pptx_file: Path):
    doc = parse(temp_pptx_file)
    assert isinstance(doc, Document)
    assert doc.metadata.file_type == "pptx"
    assert len(doc.content_tree) >= 4


def test_pptx_elements_structure(temp_pptx_file: Path):
    doc = parse(temp_pptx_file)

    # 1. Title Heading
    h1 = doc.content_tree[0]
    assert h1.type == "heading"
    assert "Architecture Review" in (h1.text or "")

    # 2. Bullet point / list item
    list_items = [e for e in doc.content_tree if e.type == "list_item"]
    assert len(list_items) >= 1
    assert "Zero GPU" in (list_items[0].text or "")

    # 3. Table
    tables = [e for e in doc.content_tree if e.type == "table"]
    assert len(tables) == 1
    assert tables[0].data is not None
    assert tables[0].data.headers == ["Metric", "Target"]
